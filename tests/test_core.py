import base64,calendar,copy,io,json,zipfile
from datetime import date
from pathlib import Path
import pytest
from openpyxl import Workbook
from docx import Document
from backend.core import *
from backend.documents import document,KINDS

def xlsx(rows,sheets=None):
 w=Workbook();s=w.active;s.title='Report'
 for r in rows:s.append(r)
 for name,rr in (sheets or {}).items():
  s=w.create_sheet(name)
  for r in rr:s.append(r)
 out=io.BytesIO();w.save(out);return out.getvalue()

def fixture(month='2026-09'):
 months=[month,shift(month,-1),shift(month,-2)]
 vals={'Sales':[10000,10000,10000],'Cost of Goods Sold':[3000,3000,3000],'Advertising':[2000,2000,2000],'Rent':[1000,1000,1000],'Other Income':[0,0,0]}
 pl=[['Profit and Loss'],['Saint Eyewear Pty Ltd'],['Exported', '2026-10-01'],['Account']+months+['Year to date']]
 for label,values in vals.items():pl.append([label]+values+[sum(values)])
 for label,values in {'Total Trading Income':[10000]*3,'Total Cost of Sales':[3000]*3,'Gross Profit':[7000]*3,'Net Profit':[4000]*3}.items():pl.append([label]+values+[sum(values)])
 bs=[['Balance Sheet'],['Saint Eyewear Pty Ltd'],['Account']+months,['Cash',30000,26000,22000],['Net Assets',30000,26000,22000],['Total Equity',30000,26000,22000],['Current Year Earnings',12000,8000,4000]]
 gl=[['General Ledger'],['Saint Eyewear Pty Ltd'],['Exported','1 Oct 2026'],['Date','Account','Contact','Description','Amount'],[date(2026,9,3),'Rent','Landlord','Rent payment',1000]]
 days=calendar.monthrange(2026,9)[1];daily=[['Date']+[date(2026,9,i) for i in range(1,days+1)],[],[],['Orders']+[1]*days,[],[],['Sales']+['333.33']*(days-1)+['333.43'],[],['Refunds']+[0]*days]
 drs=xlsx([],{'Daily Reporting':daily,'Productwise':[[],[],[None,'Frames','1','2','3','4','5','10000'],[None,'Total Sales',None,None,None,None,None,'10000']],'Countrywise':[[],[],[],[None,'Australia',30,10000],[None,'Total',30,10000]]})
 cfg=copy.deepcopy(DEFAULT_CONFIG);cfg['confirmed']=True
 payload={'month':month,'files':{k:base64.b64encode(v).decode() for k,v in {'pnl':xlsx(pl),'bs':xlsx(bs),'gl':xlsx(gl),'drs':drs}.items()},'config':cfg,'previous':{'vendors':{'Landlord':'1000'},'drs_totals':{'sales':'10000','refunds':'0'},'accounts':list(vals)},'issued':{shift(month,-1):dict((r[0],str(r[2])) for r in pl[4:])}}
 return payload

@pytest.fixture
def result():
 p=fixture();r=process(p);r['gl']=general_ledger(base64.b64decode(p['files']['gl']),p['month']);return r

def test_happy_path(result):
 assert result['status']=='review'
 assert not [x for x in result['flags'] if x['severity']=='Block']
 assert result['metrics']['2026-09']['net']=='4000'
 assert result['drs']['totals']['sales']=='10000.00'
 assert result['client_ready'] is False
 assert set(result['coverage'])=={f'R{i}' for i in range(1,15)}

def test_decimal_precision():assert dec('0.1')+dec('0.2')==dec('.3')
@pytest.mark.parametrize('bad',['NaN','Infinity',True,'banana'])
def test_bad_number(bad):
 with pytest.raises(ValidationError):dec(bad)

def test_no_gl_reconstruction():
 p=fixture();r=process(p);p['files']['gl']=base64.b64encode(xlsx([['General Ledger'],['Exported','2026-10-01'],['Date','Account','Vendor','Amount'],['2026-09-03','Rent','A',900000]])).decode();r2=process(p)
 assert r['metrics']==r2['metrics']

def test_wrong_month():
 p=fixture();p['month']='2026-08'
 with pytest.raises(ValidationError):process(p)

def test_invalid_archive():
 p=fixture();p['files']['pnl']=base64.b64encode(b'no').decode()
 with pytest.raises(ValidationError):process(p)

def test_drs_missing_sheet():
 p=fixture();p['files']['drs']=base64.b64encode(xlsx([['no']])).decode()
 with pytest.raises(ValidationError,match='DRS missing'):process(p)

def test_formula_no_cached_value():
 b=xlsx([['=1+1']])
 with pytest.raises(ValidationError,match='cached'):excel(b)

def test_wrong_export_day(result):
 r=copy.deepcopy(result);r['gl']['export_date']='2026-09-30'
 out=assess(r['pnl'],r['bs'],r['gl'],r['drs'],r['month'])
 assert any(f['rule']=='R3' and f['severity']=='Block' for f in out['flags'])

def test_bs_mismatch(result):
 r=copy.deepcopy(result);r['bs']['months']['2026-08']['Total Equity']='0'
 out=assess(r['pnl'],r['bs'],r['gl'],r['drs'],r['month'])
 assert any(f['rule']=='R2' and f['severity']=='Block' for f in out['flags'])

def test_ytd_mismatch(result):
 r=copy.deepcopy(result);r['bs']['months']['2026-09']['Current Year Earnings']='1'
 out=assess(r['pnl'],r['bs'],r['gl'],r['drs'],r['month'])
 assert any(f['rule']=='R2' and 'Current Year' in f['message'] for f in out['flags'])

def scenario(result,mutate,previous=None,issued=None):
 r=copy.deepcopy(result);mutate(r)
 return assess(r['pnl'],r['bs'],r['gl'],r['drs'],r['month'],copy.deepcopy(DEFAULT_CONFIG),previous,issued)

def rules(r):return {f['rule'] for f in r['flags']}

def test_r1(result):assert 'R1' in rules(scenario(result,lambda r:None,issued={'2026-08':{'Net Profit':'-53715'}}))
def test_r4(result):
 r=scenario(result,lambda r:r['pnl']['months']['2026-09'].__setitem__('Other Income','56624'))
 assert 'R4' in rules(r)
def test_r5(result):assert 'R5' in rules(scenario(result,lambda r:r['pnl']['months']['2026-09'].__setitem__('Total Trading Income','5600')))
def test_r6(result):assert 'R6' in rules(scenario(result,lambda r:r['pnl']['months']['2026-09'].__setitem__('Gross Profit','2000')))
def test_r7(result):assert 'R7' in rules(scenario(result,lambda r:r['pnl']['months']['2026-09'].__setitem__('Net Profit','-1000')))
def test_r8(result):
 r=scenario(result,lambda r:r['gl']['transactions'].append({'date':'2026-09-03','account':'Software Expense','vendor':'Redo Tech','description':'Subscription','amount':'12058','row':30}),previous={'vendors':{'Redo Tech':'40342'}})
 assert 'R8' in rules(r)
def test_r9(result):
 r=scenario(result,lambda r:r['gl']['transactions'].append({'date':'2026-09-04','account':'Suspense','vendor':'Unknown','description':'Uncoded','amount':'1473','row':31}))
 assert 'R9' in rules(r)
def test_r10(result):
 r=scenario(result,lambda r:None,previous={'accounts':['Sales']})
 assert 'R10' in rules(r)
def test_r11(result):
 r=scenario(result,lambda r:r['drs']['totals'].__setitem__('sales','12000'))
 assert 'R11' in rules(r)
def test_r12_duplicate(result):
 r=scenario(result,lambda r:r['gl']['transactions'].append({**r['gl']['transactions'][0],'date':'2026-09-05','row':7}))
 assert 'R12' in rules(r)
def test_r12_monthly(result):
 r=scenario(result,lambda r:r['gl']['transactions'].append({**r['gl']['transactions'][0],'description':'monthly total','amount':'4866','row':7}))
 assert 'R12' in rules(r)
def test_r13(result):
 r=scenario(result,lambda r:r['drs']['totals'].__setitem__('refunds','1000'),previous={'drs_totals':{'sales':'10000','refunds':'0'}})
 assert 'R13' in rules(r)
def test_r14(result):
 r=scenario(result,lambda r:r['gl']['transactions'].append({**r['gl']['transactions'][0],'amount':'11000','row':8}))
 assert 'R14' in rules(r)
def test_missing_history_explicit(result):
 r=scenario(result,lambda r:None)
 assert all(r['coverage'][k]=='not_evaluated' for k in ['R1','R8','R13'])

def test_unknown_account_blocks(result):
 def mutate(r):r['pnl']['accounts'].append('Unmapped');r['pnl']['months']['2026-09']['Unmapped']='100'
 assert any(f['rule']=='MAP' and f['severity']=='Block' for f in scenario(result,mutate)['flags'])

def test_zero_is_not_missing(result):
 def mutate(r):r['pnl']['months']['2026-09']['Total Trading Income']='0'
 r=scenario(result,mutate)
 assert r['metrics']['2026-09']['revenue']=='0'

def test_document_gate(result):
 r=copy.deepcopy(result);r['flags'].append({'id':'x','rule':'R3','severity':'Block','message':'Blocked'})
 with pytest.raises(ValidationError):document({'kind':'financial','run':{'figures':r}})

def test_review_note_gate(result):
 r=copy.deepcopy(result);r['flags'].append({'id':'x','rule':'R5','severity':'Review','message':'Question'})
 with pytest.raises(ValidationError):document({'kind':'financial','run':{'figures':r,'notes':{}}})

@pytest.mark.parametrize('kind',list(KINDS))
def test_documents_open_and_reopen(result,kind):
 run={'id':'synthetic-test','figures':result,'notes':{f['id']:'Synthetic test explanation only.' for f in result['flags']},'reviewed_by':'QA test - synthetic','reviewed_at':'2026-10-23T12:00:00Z'}
 out=document({'kind':kind,'run':run});raw=base64.b64decode(out['content']);doc=Document(io.BytesIO(raw));assert 'DRAFT' in '\n'.join(p.text for p in doc.paragraphs)
 z=zipfile.ZipFile(io.BytesIO(raw));assert z.namelist()[0]=='[Content_Types].xml'
 for name in z.namelist():
  if name.endswith('.xml'):ET.fromstring(z.read(name))
 if kind!='exceptions':assert '<w:tblHeader/>' in z.read('word/document.xml').decode()
 p=Path('qa/synthetic-output');p.mkdir(parents=True,exist_ok=True);(p/(kind+'.docx')).write_bytes(raw)

if __name__=='__main__':
 p=Path('qa/synthetic-fixtures');p.mkdir(parents=True,exist_ok=True)
 f=fixture()
 for k,b in f['files'].items():(p/(k+'.xlsx')).write_bytes(base64.b64decode(b))
 (p/'payload.json').write_text(json.dumps(f))

def test_r10_does_not_infer_new_account_from_zero_comparative(result):
 r=scenario(result,lambda r:None)
 assert r['coverage']['R10']=='not_evaluated'
 assert not any('New or newly active' in f['message'] for f in r['flags'])

def test_r10_zero_new_account_still_flagged(result):
 r=scenario(result,lambda r:None,previous={'accounts':['Sales','Cost of Goods Sold','Advertising','Rent']})
 assert any(f['rule']=='R10' and 'Other Income' in f['message'] for f in r['flags'])

@pytest.mark.parametrize('value,flagged',[('10000',False),('10001',False),('11000',True),('-11000',True)])
def test_large_round_threshold_boundary(result,value,flagged):
 r=scenario(result,lambda r:r['gl']['transactions'].append({**r['gl']['transactions'][0],'amount':value,'row':8}))
 assert ('R14' in rules(r)) is flagged

def test_drs_sales_threshold_not_triggered_at_exactly_one_percent(result):
 r=scenario(result,lambda r:r['drs']['totals'].__setitem__('sales','10100'))
 assert 'R11' not in rules(r)

def test_review_whitespace_does_not_unlock_document(result):
 r=copy.deepcopy(result);r['flags'].append({'id':'white','rule':'R5','severity':'Review','message':'Question'})
 with pytest.raises(ValidationError):document({'kind':'financial','run':{'figures':r,'notes':{'white':'   '},'reviewed_by':'test','reviewed_at':'2026-10-23'}})

def test_blank_drs_cells_are_not_formatted_as_zero():
 from backend.documents import money
 assert money(None)==''
 assert money('')==''
 assert money('0')=='0.00'

def test_bar_chart_has_zero_baseline_for_positive_values():
 from backend.documents import chart
 assert '<c:min val="0"/>' in chart(['one'],[('Revenue',['100'])])
